"""Tests for PDF, DOCX, and TXT text extraction."""

from pathlib import Path

import pymupdf
import pytest
from docx import Document as build_docx

from app.services.extraction import ExtractionError, extract_document


def write_pdf(path: Path, page_texts: list[str]) -> None:
    pdf = pymupdf.open()

    try:
        for page_text in page_texts:
            page = pdf.new_page()
            page.insert_text((72, 72), page_text)
        pdf.save(path)
    finally:
        pdf.close()


def test_extract_pdf_preserves_pages_and_metadata(tmp_path: Path) -> None:
    path = tmp_path / "report.pdf"
    write_pdf(path, ["First page text", "Second page text"])

    extracted = extract_document(path, "pdf")

    assert [page.page_number for page in extracted.pages] == [1, 2]
    assert [page.content for page in extracted.pages] == [
        "First page text",
        "Second page text",
    ]
    assert extracted.metadata["format"] == "pdf"
    assert extracted.metadata["page_count"] == 2
    assert extracted.pages[0].metadata["width"] > 0


def test_extract_docx_collects_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "requirements.docx"
    document = build_docx()
    document.add_paragraph("Project overview")
    document.add_paragraph("The assistant answers questions.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Team"
    document.save(path)

    extracted = extract_document(path, "docx")

    assert len(extracted.pages) == 1
    content = extracted.pages[0].content
    assert "Project overview" in content
    assert "Owner | Team" in content
    assert extracted.metadata["page_numbering"] == "unavailable"
    assert extracted.metadata["table_count"] == 1


def test_extract_txt_splits_form_feed_pages(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("Page one\fPage two", encoding="utf-8")

    extracted = extract_document(path, "txt")

    assert [page.page_number for page in extracted.pages] == [1, 2]
    assert [page.content for page in extracted.pages] == [
        "Page one",
        "Page two",
    ]
    assert extracted.metadata["page_numbering"] == "form_feed"
    assert extracted.metadata["encoding"] == "utf-8"


def test_extract_txt_supports_windows_encoding(tmp_path: Path) -> None:
    path = tmp_path / "legacy.txt"
    path.write_bytes("Caf\xe9 notes".encode("cp1252"))

    extracted = extract_document(path, "txt")

    assert extracted.pages[0].content == "Café notes"
    assert extracted.metadata["encoding"] == "cp1252"


def test_corrupt_pdf_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"this is not a pdf")

    with pytest.raises(ExtractionError, match="PDF"):
        extract_document(path, "pdf")


def test_pdf_without_text_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    write_pdf(path, [""])

    with pytest.raises(ExtractionError, match="No extractable text"):
        extract_document(path, "pdf")


def test_empty_docx_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    build_docx().save(path)

    with pytest.raises(ExtractionError, match="No extractable text"):
        extract_document(path, "docx")


def test_empty_txt_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n", encoding="utf-8")

    with pytest.raises(ExtractionError, match="does not contain any text"):
        extract_document(path, "txt")
