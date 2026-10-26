"""Text extraction for uploaded PDF, DOCX, and TXT files."""

from __future__ import annotations

import codecs
import logging
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pymupdf
from docx import Document as open_docx
from docx.table import Table
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    """Raised when a document cannot be converted into text."""


@dataclass(frozen=True)
class ExtractedPage:
    """Text and metadata for a single page."""

    page_number: int
    content: str
    metadata: dict[str, Any]


@dataclass(frozen=True)
class ExtractedDocument:
    """Result of extracting a whole document."""

    pages: list[ExtractedPage]
    metadata: dict[str, Any]


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None

    text = str(value).strip()
    return text or None


def _iso_date(value: date | datetime | None) -> str | None:
    return value.isoformat() if value is not None else None


def extract_document(path: Path, extension: str) -> ExtractedDocument:
    """Extract text and metadata from a stored document."""

    normalized = extension.lower().lstrip(".")

    if normalized == "pdf":
        return _extract_pdf(path)
    if normalized == "docx":
        return _extract_docx(path)
    if normalized == "txt":
        return _extract_txt(path)

    raise ExtractionError(f"Unsupported document format '.{normalized}'.")


def _extract_pdf(path: Path) -> ExtractedDocument:
    try:
        pdf = pymupdf.open(path)
    except Exception as exc:
        raise ExtractionError(
            "The PDF could not be opened. The file may be corrupt."
        ) from exc

    try:
        if pdf.needs_pass:
            raise ExtractionError(
                "Password-protected PDFs are not supported."
            )

        pages: list[ExtractedPage] = []
        text_characters = 0

        for page_number, page in enumerate(pdf, start=1):
            content = page.get_text("text").strip()
            text_characters += len(content)
            pages.append(
                ExtractedPage(
                    page_number=page_number,
                    content=content,
                    metadata={
                        "width": round(page.rect.width, 1),
                        "height": round(page.rect.height, 1),
                        "link_count": len(page.get_links() or []),
                    },
                )
            )

        if not pages:
            raise ExtractionError("The PDF does not contain any pages.")

        if text_characters == 0:
            raise ExtractionError(
                "No extractable text was found in the PDF. "
                "Scanned documents are not supported yet."
            )

        info = pdf.metadata or {}
        return ExtractedDocument(
            pages=pages,
            metadata={
                "format": "pdf",
                "page_count": len(pages),
                "text_characters": text_characters,
                "title": _optional_text(info.get("title")),
                "author": _optional_text(info.get("author")),
                "subject": _optional_text(info.get("subject")),
                "creator": _optional_text(info.get("creator")),
                "producer": _optional_text(info.get("producer")),
                "creation_date": _optional_text(info.get("creationDate")),
            },
        )
    except ExtractionError:
        raise
    except Exception as exc:
        logger.exception("Unexpected PDF extraction failure for %s", path)
        raise ExtractionError(
            "The PDF could not be read. The file may be corrupt."
        ) from exc
    finally:
        pdf.close()


def _extract_docx(path: Path) -> ExtractedDocument:
    try:
        document = open_docx(path)
    except Exception as exc:
        raise ExtractionError(
            "The DOCX could not be opened. The file may be corrupt."
        ) from exc

    lines: list[str] = []
    paragraph_count = 0
    table_count = 0

    try:
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if text:
                    lines.append(text)
                    paragraph_count += 1
            elif isinstance(block, Table):
                table_count += 1
                for row in block.rows:
                    cells = [
                        cell.text.strip().replace("\n", " ")
                        for cell in row.cells
                    ]
                    if any(cells):
                        lines.append(" | ".join(cells))
    except Exception as exc:
        logger.exception("Unexpected DOCX extraction failure for %s", path)
        raise ExtractionError(
            "The DOCX could not be read. The file may be corrupt."
        ) from exc

    content = "\n".join(lines)
    if not content.strip():
        raise ExtractionError(
            "No extractable text was found in the DOCX."
        )

    properties = document.core_properties
    page_metadata = {
        "page_numbering": "unavailable",
        "paragraph_count": paragraph_count,
        "table_count": table_count,
    }

    return ExtractedDocument(
        pages=[
            ExtractedPage(
                page_number=1,
                content=content,
                metadata=page_metadata,
            )
        ],
        metadata={
            "format": "docx",
            "page_count": 1,
            "page_numbering": "unavailable",
            "text_characters": len(content),
            "paragraph_count": paragraph_count,
            "table_count": table_count,
            "title": _optional_text(properties.title),
            "author": _optional_text(properties.author),
            "created": _iso_date(properties.created),
            "modified": _iso_date(properties.modified),
        },
    )


def _decode_text(raw: bytes) -> tuple[str, str]:
    if raw.startswith(codecs.BOM_UTF8):
        return "utf-8-sig", raw.decode("utf-8-sig")

    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16", raw.decode("utf-16")

    for encoding in ("utf-8", "cp1252"):
        try:
            return encoding, raw.decode(encoding)
        except UnicodeDecodeError:
            continue

    return "latin-1", raw.decode("latin-1")


def _extract_txt(path: Path) -> ExtractedDocument:
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise ExtractionError("The text file could not be read.") from exc

    encoding, decoded = _decode_text(raw)
    normalized = decoded.replace("\r\n", "\n").replace("\r", "\n")
    raw_pages = normalized.split("\f")

    if len(raw_pages) > 1:
        pages = [
            ExtractedPage(
                page_number=index,
                content=chunk.strip(),
                metadata={"encoding": encoding, "page_break": "form_feed"},
            )
            for index, chunk in enumerate(raw_pages, start=1)
        ]

        while len(pages) > 1 and not pages[-1].content:
            pages.pop()
    else:
        pages = [
            ExtractedPage(
                page_number=1,
                content=normalized.strip(),
                metadata={"encoding": encoding},
            )
        ]

    text_characters = sum(len(page.content) for page in pages)
    if text_characters == 0:
        raise ExtractionError("The text file does not contain any text.")

    return ExtractedDocument(
        pages=pages,
        metadata={
            "format": "txt",
            "page_count": len(pages),
            "page_numbering": (
                "form_feed" if len(pages) > 1 else "single_page"
            ),
            "encoding": encoding,
            "line_count": len(normalized.splitlines()),
            "text_characters": text_characters,
        },
    )
